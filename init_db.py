"""
DBを初期化し、Excel(ワールド2025.xlsx)から商品マスターを取り込む。
既存データがある場合はスキップ（--forceで再作成）。
"""
import os, sys
import psycopg2
import psycopg2.extras
from dotenv import load_dotenv

load_dotenv()
DATABASE_URL = os.environ.get('DATABASE_URL')

EXCEL_PATH = os.path.expanduser(
    '~/Library/CloudStorage/GoogleDrive-ino.inomata@gmail.com'
    '/マイドライブ/claude/ワールド2025.xlsx'
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS products (
    id          SERIAL PRIMARY KEY,
    col_group   INTEGER NOT NULL,
    category    TEXT    NOT NULL DEFAULT '',
    name        TEXT    NOT NULL,
    base_qty    FLOAT   DEFAULT 0,
    note        TEXT,
    supplier    TEXT    DEFAULT 'ワールド美容',
    sort_order  INTEGER DEFAULT 0,
    is_active   BOOLEAN DEFAULT true
);

CREATE TABLE IF NOT EXISTS stock (
    product_id  INTEGER NOT NULL,
    store       TEXT    NOT NULL DEFAULT 'イノ',
    current_qty FLOAT   DEFAULT 0,
    updated_at  TEXT,
    PRIMARY KEY (product_id, store),
    FOREIGN KEY (product_id) REFERENCES products(id)
);

CREATE TABLE IF NOT EXISTS orders (
    id          SERIAL PRIMARY KEY,
    store       TEXT    NOT NULL DEFAULT 'イノ',
    supplier    TEXT    DEFAULT 'ワールド美容',
    order_month INTEGER,
    order_day   INTEGER,
    status      TEXT    DEFAULT 'draft',
    notes       TEXT,
    created_at  TEXT
);

CREATE TABLE IF NOT EXISTS order_items (
    id          SERIAL PRIMARY KEY,
    order_id    INTEGER NOT NULL,
    product_id  INTEGER NOT NULL,
    base_qty    FLOAT,
    stock_qty   FLOAT,
    order_qty   FLOAT   NOT NULL,
    FOREIGN KEY (order_id)   REFERENCES orders(id),
    FOREIGN KEY (product_id) REFERENCES products(id)
);
"""


def import_excel(cur, excel_path):
    try:
        import openpyxl
    except ImportError:
        print("[ERROR] openpyxl がインストールされていません")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['原本']

    sort_order = 0
    current_cat = ['', '', '', '', '']

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        for grp in range(5):
            c = grp * 4
            name  = row[c]
            base  = row[c + 1]
            note_val = row[c + 2]

            if name is None:
                continue
            name = str(name).strip()
            if not name:
                continue

            if isinstance(base, (int, float)):
                note = str(note_val).strip() if isinstance(note_val, str) else None
                cur.execute(
                    "INSERT INTO products (col_group, category, name, base_qty, note, sort_order) "
                    "VALUES (%s,%s,%s,%s,%s,%s)",
                    (grp + 1, current_cat[grp], name, float(base), note, sort_order)
                )
                sort_order += 1
            else:
                if name.startswith('※') or name == 'アジアン予備欄' or name == '追加分':
                    continue
                current_cat[grp] = name

    cur.execute("SELECT COUNT(*) FROM products")
    count = cur.fetchone()[0]
    print(f"  商品 {count} 件を登録しました")


def main():
    force = '--force' in sys.argv

    conn = psycopg2.connect(DATABASE_URL)
    cur = conn.cursor()

    # スキーマ作成
    for stmt in SCHEMA.strip().split('\n\n'):
        stmt = stmt.strip()
        if stmt:
            cur.execute(stmt)
    conn.commit()
    print("スキーマを作成しました。")

    # 既存データ確認
    cur.execute("SELECT COUNT(*) FROM products")
    count = cur.fetchone()[0]

    if count > 0 and not force:
        print(f"DB は既に存在します（商品 {count} 件）。再作成するには --force を付けてください。")
        conn.close()
        return

    if count > 0 and force:
        cur.execute("DELETE FROM order_items")
        cur.execute("DELETE FROM orders")
        cur.execute("DELETE FROM stock")
        cur.execute("DELETE FROM products")
        conn.commit()
        print("既存データを削除しました。")

    if os.path.exists(EXCEL_PATH):
        print(f"Excel を読み込み中: {EXCEL_PATH}")
        import_excel(cur, EXCEL_PATH)
        conn.commit()
    else:
        print(f"[WARNING] Excel が見つかりません: {EXCEL_PATH}")
        print("  商品マスターは /products 画面から手動で追加できます。")

    conn.close()
    print("初期化完了。")


if __name__ == '__main__':
    main()
